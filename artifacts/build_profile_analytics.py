from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT_DIR = Path(__file__).resolve().parent
DOCX_PATH = OUT_DIR / "analytics_profile.docx"


TITLE = "Аналитика по профилю пользователя"
INTRO_1 = (
    "Редакция документа: подготовлена на основе текущей фактической реализации "
    "frontend, backend и связанных пользовательских сценариев профиля."
)
INTRO_2 = (
    "Назначение документа: зафиксировать корректные требования к странице профиля, "
    "операциям просмотра и редактирования данных, переходам в связанные разделы и "
    "сценарию удаления аккаунта, а также определить критерии приемки без смешения "
    "с техническими статусами задач."
)

SECTIONS = [
    (
        "1. Область действия и цель",
        [
            "Документ описывает пользовательский сценарий работы со страницей профиля "
            "в маркетплейсе после авторизации.",
            "Документ охватывает получение данных профиля через backend, отображение "
            "основной информации пользователя, редактирование ФИО, переходы в связанные "
            "разделы и удаление аккаунта.",
            "Документ не является журналом разработки и не фиксирует статусы выполнения "
            "задач. Требования и критерии приемки описаны независимо от текущего состояния кода.",
        ],
    ),
]

USE_CASE_META = [
    ("Название", "Просмотр и управление профилем пользователя."),
    ("Актор", "Авторизованный пользователь маркетплейса."),
    ("Триггер", "Пользователь открывает страницу /profile через меню в шапке сайта."),
    ("Цель", "Получить доступ к личным данным, связанным разделам и настройкам аккаунта."),
]

PRECONDITIONS = [
    "Пользователь успешно авторизован в системе.",
    "Frontend имеет доступ к API профиля по маршруту /profile/me.",
    "Backend может определить текущего пользователя через CurrentUserService.",
    "Для пользователя в системе существует запись пользователя, а профиль может быть создан автоматически при первом обращении.",
]

MAIN_FLOW = [
    "Пользователь открывает страницу /profile.",
    "Frontend отправляет запрос GET /profile/me для получения данных профиля текущего пользователя.",
    "Backend определяет пользователя, находит или создает профиль и возвращает модель ProfileResponse.",
    "Frontend отображает аватар с инициалами, полное имя пользователя, логин или email и карточки связанных разделов.",
    "Пользователь может открыть блок Настройки аккаунта.",
    "Внутри блока Настройки аккаунта пользователь может перейти к редактированию ФИО.",
    "После сохранения frontend отправляет PUT /profile/me с новыми значениями lastName, firstName и middleName.",
    "Backend сохраняет изменения и возвращает обновленные данные профиля.",
    "Пользователь может перейти из профиля в избранное по кликабельной карточке.",
    "Пользователь может открыть нижнюю часть блока Настройки аккаунта и запустить сценарий удаления аккаунта.",
]

ALTERNATIVES = [
    "Профиль отсутствует: backend автоматически создает пустой профиль при первом открытии страницы.",
    "Ошибка загрузки профиля: frontend отображает сообщение об ошибке без падения страницы.",
    "Пользователь отменяет редактирование: форма закрывается без сохранения данных.",
    "Пользователь отменяет удаление аккаунта: аккаунт и связанные данные остаются без изменений.",
    "При удалении аккаунта backend предварительно очищает связанные данные пользователя, включая корзину, избранное, SSO-связи, профиль и другие зависимые сущности.",
]

MODEL_TEXTS = [
    "Frontend работает с профилем через модуль src/api/profile.ts и страницу src/app/profile/page.tsx.",
    "Backend предоставляет API профиля через ProfileController и ProfileService.",
    "Основная модель ответа профиля содержит id, lastName, firstName, middleName, login и email.",
    "В текущей реализации email на backend возвращается из логина пользователя, если отдельное поле email отсутствует в доменной модели.",
    "Переходы из профиля в связанные разделы выполняются на уровне frontend через обычную навигацию Next.js.",
    "Удаление аккаунта является необратимым действием и должно выполняться только после явного подтверждения пользователя.",
]

REQUIREMENTS = [
    ("REQ-01", "Доступ к профилю", "Страница профиля должна быть доступна только авторизованному пользователю."),
    ("REQ-02", "Загрузка профиля", "При открытии страницы frontend должен запрашивать данные через GET /profile/me."),
    ("REQ-03", "Автосоздание профиля", "Если профиль пользователя отсутствует, backend должен создать пустую запись профиля и вернуть ее в ответе."),
    ("REQ-04", "Отображение данных", "Страница профиля должна отображать полное имя, логин или email и визуальный блок профиля."),
    ("REQ-05", "Редактирование профиля", "Пользователь должен иметь возможность изменить фамилию, имя и отчество через PUT /profile/me."),
    ("REQ-06", "Переход в избранное", "Карточка Избранное на странице профиля должна быть кликабельной и вести на /wishlist."),
    ("REQ-07", "Настройки аккаунта", "Блок удаления аккаунта должен отображаться только внутри раскрытого раздела Настройки аккаунта и располагаться в его нижней части."),
    ("REQ-08", "Удаление аккаунта", "Удаление аккаунта должно запускаться только после явного подтверждения пользователя."),
    ("REQ-09", "Очистка связанных данных", "Перед удалением пользователя backend должен очищать все связанные данные, нарушающие внешние ключи."),
    ("REQ-10", "Обработка ошибок", "Ошибки загрузки, сохранения и удаления должны отображаться пользователю в понятном виде."),
]

CRITERIA = [
    (
        "АС-01. Открытие профиля",
        [
            "Given пользователь авторизован в системе.",
            "When пользователь открывает страницу /profile.",
            "Then frontend должен запросить GET /profile/me и отобразить данные текущего пользователя.",
        ],
    ),
    (
        "АС-02. Автосоздание пустого профиля",
        [
            "Given у пользователя еще нет записи в таблице профиля.",
            "When пользователь впервые открывает страницу /profile.",
            "Then backend должен создать пустой профиль и вернуть его без ошибки 404.",
        ],
    ),
    (
        "АС-03. Редактирование ФИО",
        [
            "Given пользователь находится на странице профиля.",
            "When пользователь изменяет фамилию, имя или отчество и сохраняет форму.",
            "Then frontend должен отправить PUT /profile/me, а на странице должны отобразиться обновленные данные.",
        ],
    ),
    (
        "АС-04. Переход в избранное",
        [
            "Given пользователь находится на странице профиля.",
            "When пользователь нажимает на карточку Избранное.",
            "Then приложение должно перейти на страницу /wishlist.",
        ],
    ),
    (
        "АС-05. Видимость удаления аккаунта",
        [
            "Given пользователь находится на странице профиля.",
            "When раздел Настройки аккаунта свернут.",
            "Then блок Удалить аккаунт не должен отображаться на странице.",
            "When пользователь раскрывает раздел Настройки аккаунта.",
            "Then блок Удалить аккаунт должен появиться в нижней части этого раздела.",
        ],
    ),
    (
        "АС-06. Удаление аккаунта",
        [
            "Given пользователь подтвердил удаление аккаунта.",
            "When frontend отправляет DELETE /profile/me.",
            "Then backend должен удалить пользователя и все необходимые связанные данные без нарушения внешних ключей.",
        ],
    ),
]

COMPONENTS = [
    ("ProfileController", "Предоставляет маршруты GET /profile/me, PUT /profile/me и DELETE /profile/me для текущего пользователя."),
    ("ProfileService", "Создает профиль при необходимости, обновляет данные профиля и выполняет удаление аккаунта вместе с очисткой связанных сущностей."),
    ("ProfileRepository", "Работает с сущностью профиля и предоставляет поиск профиля по userId."),
    ("CurrentUserService", "Определяет идентификатор текущего пользователя из контекста безопасности."),
    ("profileApi", "Frontend-слой для загрузки, обновления и удаления профиля через общий apiFetch."),
    ("Profile page", "React-страница, отображающая данные пользователя, настройки аккаунта и переходы в связанные разделы."),
]

CORRECTIONS = [
    "Требования отделены от статусов реализации и сформулированы как ожидаемое поведение системы.",
    "Отдельно зафиксирован пользовательский сценарий перехода в избранное из профиля.",
    "Отдельно зафиксировано требование по видимости блока удаления аккаунта только внутри Настроек аккаунта.",
    "Добавлено требование по предварительной очистке связанных данных пользователя перед удалением аккаунта.",
]


def shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text: str, bold: bool = False) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.name = "Arial"
    run.font.size = Pt(10.5)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def build_docx() -> None:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.15

    for style_name, size, color in [
        ("Title", 24, RGBColor(0, 0, 0)),
        ("Heading 1", 16, RGBColor(0, 0, 0)),
        ("Heading 2", 13, RGBColor(0, 0, 0)),
        ("Heading 3", 11, RGBColor(67, 67, 67)),
    ]:
        style = doc.styles[style_name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.color.rgb = color
        style.paragraph_format.space_after = Pt(6)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title.add_run(TITLE)

    paragraph = doc.add_paragraph()
    paragraph.add_run("Редакция документа: ").bold = True
    paragraph.add_run(INTRO_1.split(": ", 1)[1])

    paragraph = doc.add_paragraph()
    paragraph.add_run("Назначение документа: ").bold = True
    paragraph.add_run(INTRO_2.split(": ", 1)[1])

    for heading, paragraphs in SECTIONS:
        doc.add_heading(heading, level=1)
        for item in paragraphs:
            doc.add_paragraph(item)

    doc.add_heading("2. Пользовательский сценарий (Use Case)", level=1)
    for label, value in USE_CASE_META:
        paragraph = doc.add_paragraph()
        paragraph.add_run(f"{label}: ").bold = True
        paragraph.add_run(value)

    paragraph = doc.add_paragraph()
    paragraph.add_run("Предусловия:").bold = True
    for item in PRECONDITIONS:
        doc.add_paragraph(item, style="List Bullet")

    paragraph = doc.add_paragraph()
    paragraph.add_run("Основной сценарий:").bold = True
    for item in MAIN_FLOW:
        doc.add_paragraph(item, style="List Number")

    paragraph = doc.add_paragraph()
    paragraph.add_run("Альтернативные сценарии:").bold = True
    for item in ALTERNATIVES:
        doc.add_paragraph(item, style="List Bullet")

    doc.add_heading("3. Корректная предметная и техническая модель", level=1)
    for item in MODEL_TEXTS:
        doc.add_paragraph(item)

    doc.add_heading("4. Требования", level=1)
    doc.add_paragraph(
        "Требования ниже фиксируют ожидаемое поведение системы и не содержат статусов реализации."
    )
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for index, header in enumerate(["ID", "Требование", "Описание"]):
        set_cell_text(table.rows[0].cells[index], header, bold=True)
        shade_cell(table.rows[0].cells[index], "D9E2F3")

    for rid, name, desc in REQUIREMENTS:
        row = table.add_row().cells
        for index, text in enumerate((rid, name, desc)):
            set_cell_text(row[index], text)

    doc.add_heading("5. Критерии приемки", level=1)
    doc.add_paragraph(
        "Критерии приемки описаны отдельно от требований и должны использоваться для функциональной проверки решения."
    )
    for title_text, bullets in CRITERIA:
        doc.add_paragraph(title_text, style="Heading 2")
        for bullet in bullets:
            doc.add_paragraph(bullet, style="List Bullet")

    doc.add_heading("6. Техническое описание компонентов", level=1)
    for name, desc in COMPONENTS:
        paragraph = doc.add_paragraph()
        paragraph.add_run(f"{name}: ").bold = True
        paragraph.add_run(desc)

    doc.add_heading("7. Исправления относительно исходного документа", level=1)
    for item in CORRECTIONS:
        doc.add_paragraph(item, style="List Bullet")

    doc.save(DOCX_PATH)


if __name__ == "__main__":
    build_docx()
    print(DOCX_PATH)
